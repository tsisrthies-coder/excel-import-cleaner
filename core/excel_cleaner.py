import os
import openpyxl
from openpyxl.cell import Cell
from .date_converter import is_text_date, format_text_date, convert_to_string_date

class ExcelCleaner:
    def __init__(self, filepath, sheets_to_process, options, log_cb, progress_cb):
        self.filepath = filepath
        self.sheets_to_process = sheets_to_process
        self.options = options
        self.log = log_cb
        self.progress = progress_cb
        
        self.stats = {
            'processed_cells': 0,
            'dates_converted': 0,
            'cells_cleaned': 0,
            'values_to_text': 0
        }

    def process(self, output_path):
        self.log("Démarrage du chargement du fichier (cela peut prendre un instant)...")
        ext = os.path.splitext(self.filepath)[1].lower()
        keep_vba = True if ext == '.xlsm' else False
        
        # Chargement complet (data_only=False pour préserver les formules)
        wb = openpyxl.load_workbook(self.filepath, keep_vba=keep_vba)
        self.log(f"Fichier chargé. {len(self.sheets_to_process)} feuille(s) à traiter.")

        total_cells_to_process = 0
        for sheet_name in self.sheets_to_process:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                total_cells_to_process += ws.max_row * ws.max_column

        current_cell = 0

        for sheet_name in self.sheets_to_process:
            if sheet_name not in wb.sheetnames:
                continue
                
            self.log(f"Traitement de la feuille : {sheet_name}")
            ws = wb[sheet_name]
            
            for row in ws.iter_rows():
                for cell in row:
                    current_cell += 1
                    if current_cell % 1000 == 0:
                        self.progress(current_cell / total_cells_to_process, self.stats)

                    if cell.value is None:
                        continue
                        
                    # Ignorer les formules
                    if cell.data_type == 'f' or str(cell.value).startswith('='):
                        continue
                    
                    self.stats['processed_cells'] += 1
                    original_val = cell.value

                    # 1. Conversion des Vraies Dates Excel
                    if cell.is_date:
                        new_val = convert_to_string_date(original_val)
                        cell.value = new_val
                        cell.data_type = 's'
                        self.stats['dates_converted'] += 1
                        continue

                    # 2. Conversion des Dates déjà en format Texte
                    if isinstance(original_val, str) and is_text_date(original_val):
                        new_val = format_text_date(original_val)
                        if new_val != original_val:
                            cell.value = new_val
                            cell.data_type = 's'
                            self.stats['dates_converted'] += 1
                        continue

                    # 3. Nettoyage des chaînes de caractères
                    if isinstance(original_val, str):
                        # Préserve les zéros, nettoie juste les espaces superflus/invisibles
                        cleaned_val = original_val.strip(' \t\n\r')
                        if cleaned_val != original_val:
                            cell.value = cleaned_val
                            cell.data_type = 's'
                            self.stats['cells_cleaned'] += 1
                        continue

                    # 4. Conversion des nombres en texte (Mode Auto)
                    if self.options.get('auto_mode') and isinstance(original_val, (int, float)):
                        cell.value = str(original_val)
                        cell.data_type = 's'
                        self.stats['values_to_text'] += 1

        self.progress(1.0, self.stats)
        self.log("Sauvegarde du nouveau fichier en cours...")
        wb.save(output_path)
        wb.close()
        self.log("Traitement terminé avec succès !")
        return self.stats